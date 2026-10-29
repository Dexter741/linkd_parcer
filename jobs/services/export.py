"""CSV and Excel export of jobs together with their companies.

Columns are ordered for reading: the essentials first, long text last. Company
profile columns are optional: a column that is empty (or nearly empty, on the
Jobs sheet) in the exported data is left out, so a partial company run does not
produce a sheet of blank columns.
"""

from __future__ import annotations

import csv
from collections.abc import Iterable, Iterator
from dataclasses import dataclass
from typing import IO

from django.db.models import CharField, Count, OuterRef, Q, QuerySet, Subquery, TextField
from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from jobs.models import Company, Job


@dataclass(frozen=True)
class Column:
    header: str
    path: str  # attribute path; "company.x" reads the related Company
    width: int = 16
    optional: bool = False  # left out when (nearly) empty in the exported data


JOB_COLUMNS: list[Column] = [
    Column("LinkedIn ID", "linkedin_id", 13),
    Column("Title", "title", 40),
    Column("Company", "company_name", 28),
    Column("Company LinkedIn", "company_url", 30),
    Column("Location", "location", 22),
    Column("Workplace type", "workplace_type", 12, optional=True),
    Column("Listed", "listed_at", 11),
    Column("Posted", "posted_ago", 13),
    Column("Seniority", "seniority_level", 16),
    Column("Employment type", "employment_type", 14),
    Column("Job function", "job_function", 24),
    Column("Industry", "industries", 26),
    Column("Salary", "salary", 22, optional=True),
    Column("Applicants", "applicants", 14),
    Column("Closed", "is_closed", 8),
    Column("Job URL", "url", 32),
    Column("Apply URL", "apply_url", 32, optional=True),
    Column("Company website", "company.website", 26, optional=True),
    Column("Company size", "company.company_size", 18, optional=True),
    Column("Company employees on LinkedIn", "company.employees_on_linkedin", 12, optional=True),
    Column("Company headquarters", "company.headquarters", 22, optional=True),
    Column("Company address", "company.address", 30, optional=True),
    Column("Company type", "company.company_type", 16, optional=True),
    Column("Company founded", "company.founded", 10, optional=True),
    Column("Company followers", "company.followers", 12, optional=True),
    Column("Company specialties", "company.specialties", 30, optional=True),
    Column("Company tagline", "company.tagline", 30, optional=True),
    Column("Company about", "company.description", 50, optional=True),
    Column("Description", "description_text", 100),
]

COMPANY_COLUMNS: list[Column] = [
    Column("Name", "name", 32),
    Column("LinkedIn", "linkedin_url", 36),
    Column("Industry", "industry_or_jobs", 28),
    Column("Jobs collected", "job_count", 10),
    Column("Website", "website", 28, optional=True),
    Column("Size", "company_size", 18, optional=True),
    Column("Employees on LinkedIn", "employees_on_linkedin", 12, optional=True),
    Column("Headquarters", "headquarters", 22, optional=True),
    Column("Address", "address", 30, optional=True),
    Column("Type", "company_type", 16, optional=True),
    Column("Founded", "founded", 10, optional=True),
    Column("Followers", "followers", 12, optional=True),
    Column("Specialties", "specialties", 30, optional=True),
    Column("Tagline", "tagline", 30, optional=True),
    Column("About", "description", 80, optional=True),
]

# An optional column on the Jobs sheet is kept when at least this share of rows has a value.
JOB_SHEET_MIN_FILL = 0.05

EXCEL_CELL_LIMIT = 32767
FORMULA_PREFIXES = ("=", "+", "-", "@", "\t", "\r")
DISPLAY_CHOICES = {"workplace_type": Job.WorkplaceType}


def _neutralize_formula(value):
    """Scraped text starting with '=' etc. would run as a formula in Excel
    (CSV/formula injection). A leading apostrophe makes it plain text."""
    if isinstance(value, str) and value.startswith(FORMULA_PREFIXES):
        return "'" + value
    return value


def _value(obj, path: str):
    for part in path.split("."):
        if obj is None:
            return None
        obj = getattr(obj, part)
    if path in DISPLAY_CHOICES and obj is not None:
        obj = DISPLAY_CHOICES[path](obj).label if obj else ""
    if hasattr(obj, "tzinfo") and obj.tzinfo is not None:
        obj = obj.replace(tzinfo=None, microsecond=0)
    return _neutralize_formula(obj)


def _filled(queryset: QuerySet, path: str) -> int:
    """How many rows have a value in this column."""
    lookup = path.replace(".", "__")
    model = queryset.model
    for part in path.split(".")[:-1]:
        model = model._meta.get_field(part).related_model
    field = model._meta.get_field(path.split(".")[-1])
    q = queryset.exclude(**{f"{lookup}__isnull": True})
    if isinstance(field, CharField | TextField):
        q = q.exclude(**{lookup: ""})
    return q.count()


def _columns(queryset: QuerySet, columns: list[Column], min_fill: float) -> list[Column]:
    total = queryset.count()
    kept = []
    for column in columns:
        if not column.optional:
            kept.append(column)
            continue
        filled = _filled(queryset, column.path)
        if filled and filled >= total * min_fill:
            kept.append(column)
    return kept


def job_columns(jobs: QuerySet[Job]) -> list[Column]:
    return _columns(jobs, JOB_COLUMNS, JOB_SHEET_MIN_FILL)


def _companies(jobs: QuerySet[Job]) -> QuerySet[Company]:
    ids = jobs.exclude(company=None).values("company_id")
    industry_from_jobs = Subquery(
        Job.objects.filter(company=OuterRef("pk")).exclude(industries="").values("industries")[:1]
    )
    return (
        Company.objects.filter(pk__in=ids)
        .annotate(job_count=Count("jobs", filter=Q(jobs__in=jobs)), industry_from_jobs=industry_from_jobs)
        .order_by("-job_count", "name")
    )


def company_columns(jobs: QuerySet[Job]) -> list[Column]:
    return _columns(_companies(jobs), COMPANY_COLUMNS, 0)


def _rows(records: Iterable, columns: list[Column]) -> Iterator[list]:
    for record in records:
        yield [_value(record, c.path) for c in columns]


def job_rows(jobs: QuerySet[Job], columns: list[Column]) -> Iterator[list]:
    return _rows(jobs.select_related("company").iterator(chunk_size=500), columns)


def company_rows(jobs: QuerySet[Job], columns: list[Column]) -> Iterator[list]:
    return _rows(_companies(jobs).iterator(chunk_size=500), columns)


def csv_lines(jobs: QuerySet[Job]) -> Iterator[str]:
    """CSV as a stream of lines, for a streaming HTTP response."""

    class Echo:
        def write(self, value):
            return value

    writer = csv.writer(Echo())
    columns = job_columns(jobs)
    yield writer.writerow([c.header for c in columns])
    for row in job_rows(jobs, columns):
        yield writer.writerow(["" if v is None else v for v in row])


def write_csv(jobs: QuerySet[Job], out: IO[str]) -> int:
    """Write CSV to a text stream; returns the number of jobs. Callers should open
    files with encoding='utf-8-sig' so Excel detects UTF-8."""
    lines = csv_lines(jobs)
    out.write(next(lines))  # header
    count = 0
    for line in lines:
        out.write(line)
        count += 1
    return count


def _excel_safe(value):
    if isinstance(value, str):
        value = ILLEGAL_CHARACTERS_RE.sub("", value)
        if len(value) > EXCEL_CELL_LIMIT:
            value = value[: EXCEL_CELL_LIMIT - 1] + "…"
    return value


HEADER_FONT = Font(bold=True, color="FFFFFF")
HEADER_FILL = PatternFill("solid", fgColor="2F5597")
ROW_HEIGHT = 15


def _sheet(wb: Workbook, title: str, columns: list[Column], rows: Iterator[list]) -> int:
    ws = wb.create_sheet(title)
    ws.append([c.header for c in columns])
    for i, column in enumerate(columns, start=1):
        ws.column_dimensions[get_column_letter(i)].width = column.width
        cell = ws.cell(row=1, column=i)
        cell.font, cell.fill = HEADER_FONT, HEADER_FILL
    count = 0
    for count, row in enumerate(rows, start=1):
        ws.append([_excel_safe(v) for v in row])
        # Fixed height and no wrapping: long descriptions stay on one line in the
        # grid instead of making each row screen-high; the full text is in the cell.
        ws.row_dimensions[count + 1].height = ROW_HEIGHT
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(columns))}{count + 1}"
    return count


def write_xlsx(jobs: QuerySet[Job], out: IO[bytes] | str) -> int:
    """Two sheets: Jobs (with company columns) and Companies. Returns the number of jobs."""
    wb = Workbook()
    wb.remove(wb.active)
    columns = job_columns(jobs)
    count = _sheet(wb, "Jobs", columns, job_rows(jobs, columns))
    columns = company_columns(jobs)
    _sheet(wb, "Companies", columns, company_rows(jobs, columns))
    wb.save(out)
    return count
